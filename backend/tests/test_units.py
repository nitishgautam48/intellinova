import io

from pptx import Presentation
from pptx.util import Inches

from app.services import learner, parsing, questions
from app.services.grading import grade
from app.models import Question


def test_youtube_ids():
    assert parsing.youtube_ids("https://youtu.be/abc123") == ("abc123", None)
    assert parsing.youtube_ids("https://www.youtube.com/watch?v=abc&list=PL1") == ("abc", "PL1")
    assert parsing.youtube_ids("https://www.youtube.com/playlist?list=PL9") == (None, "PL9")
    assert not parsing.is_youtube("https://example.com/watch?v=x")


def test_transcript_grouping_keeps_timestamps():
    rows = [{"text": f"Sentence {i}.", "start": i * 10.0, "duration": 10.0} for i in range(12)]
    segs = parsing.group_transcript(rows, "Video 2 · ", window=50)
    assert segs[0].location == "Video 2 · 00:00" and segs[0].kind == "Transcript"
    assert segs[1].t_start >= 50 and segs[1].location.startswith("Video 2 · 0")


def test_pptx_slides(tmp_path, monkeypatch):
    monkeypatch.setattr("app.config.settings.storage_dir", str(tmp_path))
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Resistors in series"
    s.placeholders[1].text_frame.text = "Same current flows\nRs = R1 + R2 + R3"
    s.notes_slide.notes_text_frame.text = "Remind the class about fuses"
    tb = prs.slides.add_slide(prs.slide_layouts[5])
    tb.shapes.title.text = "Parallel"
    tb.shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(1)).text_frame.text = "1/Rp = 1/R1 + 1/R2"
    buf = io.BytesIO()
    prs.save(buf)
    p = tmp_path / "deck.pptx"
    p.write_bytes(buf.getvalue())
    out = parsing.parse_pptx(str(p))
    assert out.size_label == "2 slides"
    assert out.segments[0].location == "Slide 1" and out.segments[0].heading == "Resistors in series"
    assert "Notes: Remind" in out.segments[0].text and "1/Rp" in out.segments[1].text


def test_chunking_respects_limit():
    long = " ".join(["word."] * 500)
    chunks = parsing.chunk_paragraphs([long], max_words=100)
    assert all(len(c.split()) <= 110 for c in chunks) and len(chunks) >= 5


def test_sympy_and_numbers():
    assert questions.eval_expression("220*4") == 880
    assert questions.eval_expression("2^2*5*10") == 200
    assert questions.eval_expression("__import__('os')") is None
    assert questions.num("8.8 × 10^2 W") == 880.0
    assert questions.num("1,200 J") == 1200.0
    assert questions.num("3.6 × 10⁶ J") == 3.6e6
    assert questions.num("1.6 x 10^-8 Ω m") == 1.6e-8
    assert questions.num("−4.5") == -4.5


def test_grading_mcq_numeric_misconception():
    q = Question(qtype="Numerical", answer="880", answer_num=880.0, tolerance=0.01, misconceptions={"55.0": "Divides V by I"},
                 rubric=[], options=[])
    assert grade(q, "880 W")["grade"] == "correct"
    assert grade(q, "879.5")["grade"] == "correct"
    g = grade(q, "55")
    assert g["grade"] == "wrong" and g["misconception"] == "Divides V by I"
    m = Question(qtype="MCQ", answer="3 Ω", options=["12 Ω", "6 Ω", "3 Ω", "36 Ω"], misconceptions={"12 Ω": "Adds in parallel"}, rubric=[])
    assert grade(m, "3 Ω")["credit"] == 1.0 and grade(m, "12 Ω")["misconception"] == "Adds in parallel"


def test_bkt_moves_the_right_way():
    p = 0.3
    up = learner.bkt_update(p, 1.0, 0.1, 0.25, 0.1)
    down = learner.bkt_update(p, 0.0, 0.1, 0.25, 0.1)
    half = learner.bkt_update(p, 0.5, 0.1, 0.25, 0.1)
    assert down < p < up and down < half < up
    # a correct numerical (low guess) is stronger evidence than a correct MCQ
    assert learner.bkt_update(p, 1.0, 0.1, 0.05, 0.1) > up


def test_target_difficulty():
    assert learner.target_difficulty(0.2) == "Easy"
    assert learner.target_difficulty(0.7) == "Medium"
    assert learner.target_difficulty(0.95) == "Hard"


def test_norm_class():
    from app.routers.common import norm_class

    assert {norm_class(x) for x in ("10", "Class 10", "class  10", "Grade 10", "10th", " std 10 ")} == {"Class 10"}
    assert norm_class("Starting college") == "Starting college"
    assert norm_class("") == ""
