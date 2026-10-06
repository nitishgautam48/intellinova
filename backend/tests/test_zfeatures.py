"""Problem-statement features added after the core flow: topics/subtopics identified from material (1c),
course flow map (6a), exam study schedule (6c), weak-topic revision packs and slides (6b),
engine-driven personalisation evaluation (5c) and voice tutoring (6e)."""
import io

import pymupdf

from tests.conftest import mail_code, new_client

F: dict = {}

MAGNET_PAGES = [
    ("Magnetic field lines", "Magnetic field lines emerge from the north pole and merge at the south pole. They never intersect."),
    ("Right hand thumb rule", "Hold a current carrying conductor in the right hand with the thumb along the current; the curled "
                              "fingers show the direction of the magnetic field around it."),
]


def ok(r, code=200):
    assert r.status_code == code, f"{r.status_code}: {r.text}"
    return r.json()


def pdf(pages) -> bytes:
    doc = pymupdf.open()
    for head, body in pages:
        page = doc.new_page()
        page.insert_text((72, 80), head, fontsize=20)
        page.insert_textbox(pymupdf.Rect(72, 110, 520, 700), body, fontsize=11)
    return doc.tobytes()


def test_new_material_proposes_topics_with_subtopics(admin):
    s = ok(admin.post("/api/admin/curriculum/subjects", json={"board": "CBSE", "class_level": "Class 10", "name": "Physics"}))
    F["subject"] = s["id"]
    files = {"file": ("magnetism.pdf", io.BytesIO(pdf(MAGNET_PAGES)), "application/pdf")}
    src = ok(admin.post("/api/admin/kb/sources", data={"kind": "textbook", "title": "Magnetism notes", "subject_id": s["id"]}, files=files))
    kb = ok(admin.get("/api/admin/kb"))
    row = next(x for x in kb["sources"] if x["id"] == src["id"])
    assert row["stage"] == 6 and row["proposals"] == 1 and kb["pending_proposals"] >= 1
    assert "Identify topics" in kb["stages"]

    props = ok(admin.get("/api/admin/kb/proposals"))
    p = next(x for x in props if x["source_id"] == src["id"])
    assert p["new_chapter"] and p["chapter"] == "Magnetic Effects of Current" and p["units"] >= 2
    assert p["subtopics"][:2] == ["Part 1", "Part 2"] and p["samples"]

    acc = ok(admin.post(f"/api/admin/kb/proposals/{p['id']}/accept", json={}))
    assert acc["status"] == "accepted"
    tree = ok(admin.get(f"/api/admin/curriculum/subjects/{s['id']}"))
    ch = next(c for c in tree["chapters"] if c["name"] == "Magnetic Effects of Current")
    topic = ch["topics"][0]
    assert not topic["published"] and topic["subtopics"] == ["Part 1", "Part 2"] and tree["draft"]
    units = ok(admin.get("/api/admin/kb/units", params={"source_id": src["id"]}))["units"]
    assert all(u["topic_id"] == topic["id"] and not u["low"] for u in units if u["kind"] == "Text")
    assert {u["subtopic"] for u in units if u["kind"] == "Text"} == {"Part 1", "Part 2"}
    assert ok(admin.get("/api/admin/kb/proposals")) == [x for x in ok(admin.get("/api/admin/kb/proposals")) if x["source_id"] != src["id"]]
    r = admin.post(f"/api/admin/kb/proposals/{p['id']}/accept", json={})
    assert r.status_code == 400  # already handled

    # Editing a unit's subtopic by hand
    u = next(u for u in units if u["kind"] == "Text")
    upd = ok(admin.patch(f"/api/admin/kb/units/{u['id']}", json={"subtopic": "  Field   lines "}))
    assert upd["subtopic"] == "Field lines"
    F["topic"], F["chapter"] = topic["id"], ch["id"]
    ok(admin.post(f"/api/admin/curriculum/subjects/{s['id']}/publish", json={}))


def test_students_see_subtopics(admin, stamp):
    c = new_client()
    email = f"arjun{stamp}@example.com"
    ok(c.post("/api/auth/signup", json={"name": "Arjun Rao", "email": email, "password": "studentpw1", "accepted_terms": True}))
    ok(c.post("/api/auth/verify", json={"type": "signup", "email": email, "token": mail_code(email)}))
    ok(c.put("/api/me/profile", json={"board": "CBSE", "class_level": "Class 10", "languages": ["English"],
                                     "subject_ids": [F["subject"]], "diag_mode": "skip"}))
    ok(c.post("/api/me/onboarding/complete"))
    F["student"] = c
    chap = ok(c.get(f"/api/learn/chapters/{F['chapter']}"))
    assert set(chap["topics"][0]["subtopics"]) == {"Field lines", "Part 2"}


def test_flow_map(admin):
    sid, t1 = F["subject"], F["topic"]
    tree = ok(admin.post(f"/api/admin/curriculum/chapters/{F['chapter']}/topics", json={"name": "Force on a conductor"}))
    t2 = next(t for c in tree["chapters"] if c["id"] == F["chapter"] for t in c["topics"] if t["name"] == "Force on a conductor")["id"]
    ok(admin.patch(f"/api/admin/curriculum/topics/{t2}", json={"prereq_ids": [t1]}))

    # Admins see drafts and the new prerequisite link before publishing; students don't.
    m = ok(admin.get(f"/api/admin/curriculum/subjects/{sid}/map"))
    nodes = {n["id"]: n for n in m["nodes"]}
    assert not nodes[t2]["published"] and nodes[t2]["level"] == 1 and nodes[t1]["level"] == 0
    assert {"from": t1, "to": t2} in m["edges"] and nodes[t1]["subtopics"]
    c = F["student"]
    assert t2 not in {n["id"] for n in ok(c.get(f"/api/learn/subjects/{sid}/map"))["nodes"]}

    ok(admin.post(f"/api/admin/curriculum/subjects/{sid}/publish", json={}))
    m = ok(c.get(f"/api/learn/subjects/{sid}/map"))
    nodes = {n["id"]: n for n in m["nodes"]}
    assert nodes[t1]["state"] == "current" and nodes[t2]["state"] == "todo" and nodes[t1]["status"] == "Not assessed"
    assert [x["name"] for x in m["chapters"]] == ["Magnetic Effects of Current"]
    ok(c.post(f"/api/learn/topics/{t1}/complete"))
    nodes = {n["id"]: n for n in ok(c.get(f"/api/learn/subjects/{sid}/map"))["nodes"]}
    assert nodes[t1]["state"] == "done" and nodes[t2]["state"] == "current"
    F["topic2"] = t2


def _user_id(c) -> str:
    return ok(c.get("/api/auth/me"))["id"]


def test_exam_schedule(admin):
    import uuid
    from datetime import date, timedelta

    from app.db import SessionLocal
    from app.services import learner

    c, sid, t1, t2 = F["student"], F["subject"], F["topic"], F["topic2"]
    uid = uuid.UUID(_user_id(c))
    with SessionLocal() as db:  # two wrong answers on the finished topic make it weak
        for _ in range(2):
            learner.update(db, uid, uuid.UUID(t1), 0.0, "quiz", evidence="Quiz · test")
        db.commit()

    assert ok(c.get("/api/schedule"))["schedule"] is None
    opts = ok(c.get("/api/schedule/options"))
    assert opts[0]["id"] == sid and opts[0]["chapters"][0]["id"] == F["chapter"]
    assert c.put("/api/schedule", json={"exam_date": date.today().isoformat(), "subject_ids": [sid]}).status_code == 400
    assert c.put("/api/schedule", json={"exam_date": (date.today() + timedelta(days=5)).isoformat(), "subject_ids": [sid],
                                        "rest_days": [0, 1, 2, 3, 4, 5]}).status_code == 400

    exam = date.today() + timedelta(days=14)
    rest = (date.today().weekday() + 1) % 7  # a weekly rest day that isn't today, so today has work
    s = ok(c.put("/api/schedule", json={"title": "Physics unit test", "exam_date": exam.isoformat(), "minutes_per_day": 45,
                                        "subject_ids": [sid], "rest_days": [rest]}))["schedule"]
    assert s["days_left"] == 14 and s["summary"]["to_learn"] == 1 and s["summary"]["to_fix"] == 1 and s["summary"]["fits"]
    assert all(date.fromisoformat(d["date"]).weekday() != rest for d in s["days"])
    assert s["days"][-1]["date"] < exam.isoformat() and s["days"][-1]["final"]
    items = [(d["date"], it) for d in s["days"] for it in d["items"]]
    learn = next(x for x in items if x[1]["kind"] == "learn")
    assert learn[1]["topic_id"] == t2
    # practice the day after learning, spaced reviews later, the weak topic fixed early, a mock test at the end
    assert any(it["kind"] == "practice" and it["topic_id"] == t2 and d > learn[0] for d, it in items)
    assert any(it["kind"] == "review" and it["topic_id"] == t2 and not dd["final"] for dd in s["days"] for it in dd["items"])
    weak = next(it for d, it in items if it["kind"] == "practice" and it["topic_id"] == t1)
    assert weak["why"].startswith("Weak")
    assert s["days"][-1]["items"][0]["kind"] == "mock"
    assert all(d["minutes"] <= 45 or len(d["items"]) == 1 for d in s["days"])
    f = next(x for x in s["forgetting"] if x["topic_id"] == t1)
    assert f["recall_exam"] <= f["recall_now"] and f["status"] == "Weak"

    # Today's plan on the dashboard follows the schedule
    dash = ok(c.get("/api/dashboard"))
    assert dash["plan"]["title"] == "Physics unit test" and dash["plan"]["days_left"] == 14
    today = s["today"]
    assert today and [t["title"] for t in dash["tasks"]][0].endswith(today[0]["name"])
    assert "Physics unit test in 14 days" == dash["tasks"][0]["why"]

    # Ticking an item, and doing the work, both complete it
    key = today[0]["key"]
    assert ok(c.post(f"/api/schedule/items/{key}/toggle"))["done"]
    s = ok(c.get("/api/schedule"))["schedule"]
    assert s["today"][0]["done"] and s["done"] >= 1
    if any(it["kind"] == "learn" for it in s["today"]):
        ok(c.post(f"/api/learn/topics/{t2}/complete"))
        assert next(it for it in ok(c.get("/api/schedule"))["schedule"]["today"] if it["kind"] == "learn")["done"]

    # Rebuilding re-plans from the learner model: the learned topic drops out of "to learn"
    s = ok(c.post("/api/schedule/rebuild"))["schedule"]
    assert s["summary"]["to_learn"] == 0
    ok(c.delete("/api/schedule"))
    assert ok(c.get("/api/schedule"))["schedule"] is None
    assert ok(c.get("/api/dashboard"))["plan"] is None


def test_revision_pack_and_slides(admin):
    import pptx

    c, t1, t2 = F["student"], F["topic"], F["topic2"]
    r = c.post("/api/study/revision-pack", json={"topic_ids": [t2]})
    assert r.status_code == 400 and "no class material" in r.json()["detail"].lower()

    # One click with no topics picks the weakest ones from the learner model.
    p = ok(c.post("/api/study/revision-pack", json={}))
    assert p["topics"] == [ok(c.get(f"/api/learn/chapters/{F['chapter']}"))["topics"][0]["name"]]
    m = ok(c.get(f"/api/study/materials/{p['id']}"))
    assert m["status"] == "Ready" and m["kind"] == "pack" and m["topic_ids"] == [t1] and m["t"].startswith("Revision pack:")
    assert m["outputs"]["full"] and m["flashcards"] and m["outputs"]["audio_script"]
    assert all("Magnetism notes" in s["src"]["loc"] for s in m["outputs"]["full"])
    lib = ok(c.get("/api/study/materials"))
    assert next(x for x in lib if x["id"] == p["id"])["s"].endswith("Revision pack")

    r = c.get(f"/api/study/materials/{p['id']}/slides")
    assert r.status_code == 200 and r.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.presentationml")
    deck = pptx.Presentation(io.BytesIO(r.content))
    texts = [" ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame) for s in deck.slides]
    assert m["t"] in texts[0] and any("Formulas" in t for t in texts) and any("Check yourself" in t for t in texts)
    assert len(deck.slides) >= 2 + len(m["outputs"]["full"])


def test_simulated_students_use_the_real_quiz_engine(admin):
    import uuid

    from sqlalchemy import func, select

    from app.db import SessionLocal
    from app.models import Mastery, Question, Quiz, User
    from app.services import evaluation

    t1 = uuid.UUID(F["topic"])
    with SessionLocal() as db:
        for i in range(12):
            d = ("Easy", "Medium", "Hard")[i % 3]
            if i % 2:
                db.add(Question(topic_id=t1, qtype="Numerical", difficulty=d, stem=f"Sim numerical {i}?", answer=str(i + 1),
                                answer_num=float(i + 1), status="verified"))
            else:
                db.add(Question(topic_id=t1, qtype="MCQ", difficulty=d, stem=f"Sim MCQ {i}?", options=["A", "B", "C", "D"],
                                answer="B", status="verified"))
        db.commit()
        before = {q.id: (q.irt_b, q.times_served) for q in db.scalars(select(Question).where(Question.topic_id == t1))}
        users0 = db.scalar(select(func.count()).select_from(User))
        quizzes0 = db.scalar(select(func.count()).select_from(Quiz))
        mast0 = db.scalar(select(func.count()).select_from(Mastery))

        sim = evaluation.run_engine_simulation(db, sessions=4, students=1)
        assert sim and sim["mode"] == "engine" and sim["topics"] and sim["answers"] > 0
        assert len(sim["adaptive"]) == 5 and len(sim["personas"]) == len(evaluation.PERSONAS)
        assert sim["adaptive"][0] == sim["baseline"][0] == evaluation.learner.P_INIT
        assert sim["adaptive"][-1] != sim["adaptive"][0]  # the engine's BKT updates moved the estimates

        # Nothing the simulated students did was kept.
        db.expire_all()
        assert db.scalar(select(func.count()).select_from(User)) == users0
        assert db.scalar(select(func.count()).select_from(Quiz)) == quizzes0
        assert db.scalar(select(func.count()).select_from(Mastery)) == mast0
        after = {q.id: (q.irt_b, q.times_served) for q in db.scalars(select(Question).where(Question.topic_id == t1))}
        assert after == before
        assert not db.scalar(select(func.count()).select_from(User).where(User.email.like("sim-%@simulation.invalid")))


def test_voice_questions_are_transcribed_on_the_server(monkeypatch):
    from app.services import stt

    c = F["student"]
    seen = {}

    def fake(data, lang, suffix):
        seen.update(n=len(data), lang=lang, suffix=suffix)
        return {"text": "What is the right hand thumb rule?", "language": "en", "duration": 2.4}

    monkeypatch.setattr(stt, "transcribe", fake)
    monkeypatch.setattr(stt, "available", lambda: True)
    assert ok(c.get("/api/tutor/voice", params={"lang": "hi"}))["stt"] is True

    audio = b"\x1aE\xdf\xa3" + b"\x00" * 4000  # webm header + payload
    r = ok(c.post("/api/tutor/transcribe", data={"lang": "hing"}, files={"audio": ("q.webm", audio, "audio/webm;codecs=opus")}))
    assert r["text"].startswith("What is") and seen == {"n": len(audio), "lang": "hing", "suffix": ".webm"}
    assert ok(c.post("/api/tutor/transcribe", files={"audio": ("q.webm", b"x" * 100, "audio/webm")}))["text"] == ""
    assert c.post("/api/tutor/transcribe", files={"audio": ("q.txt", audio, "text/plain")}).status_code == 415
    assert c.post("/api/tutor/transcribe", files={"audio": ("q.webm", b"x" * (stt.MAX_BYTES + 10), "audio/webm")}).status_code == 413
    monkeypatch.setattr(stt, "available", lambda: False)
    assert c.post("/api/tutor/transcribe", files={"audio": ("q.webm", audio, "audio/webm")}).status_code == 501


def test_session_cookies_let_pages_survive_access_expiry(stamp):
    c = new_client()
    email = f"kavya{stamp}@example.com"
    ok(c.post("/api/auth/signup", json={"name": "Kavya Nair", "email": email, "password": "studentpw1", "accepted_terms": True}))
    r = c.post("/api/auth/verify", json={"type": "signup", "email": email, "token": mail_code(email)})
    ok(r)
    assert {"inn_at", "inn_rt", "inn_s"} <= set(c.cookies.keys())
    # The access cookie expires first: the refresh token (sent only to /api/auth) renews the session.
    c.cookies.delete("inn_at")
    assert c.get("/api/me/header").status_code == 401
    ok(c.post("/api/auth/refresh"))
    assert ok(c.get("/api/me/header"))["name"]
    # A dead refresh token clears everything, so the page gate sends the student to log in once (no loop).
    c.cookies.set("inn_rt", "not-a-token", path="/api/auth")
    c.cookies.delete("inn_at")
    assert c.post("/api/auth/refresh").status_code == 401
    assert "inn_s" not in c.cookies and "inn_at" not in c.cookies


def test_team_test_set_import_and_drafting(admin):
    t = admin.get("/api/admin/eval/cases/template.csv")
    assert t.status_code == 200 and t.text.startswith("question,expected_answer,expected_locations,category,off_material")

    csv_set = ("question,expected_answer,expected_locations,category,off_material\n"
               "What do magnetic field lines never do?,They never intersect.,p. 1,Textbook page,no\n"
               "How do you find the direction of the field around a wire?,Right hand thumb rule.,p. 2; p. 1,Textbook page,no\n"
               "What is the capital of Australia?,,,Textbook page,yes\n")
    r = ok(admin.post("/api/admin/eval/cases/import", data={"subject_id": F["subject"]},
                      files={"file": ("set.csv", csv_set.encode(), "text/csv")}))
    assert r == {"added": 3, "skipped": 0}
    r = ok(admin.post("/api/admin/eval/cases/import", files={"file": ("set.csv", csv_set.encode(), "text/csv")}))
    assert r == {"added": 0, "skipped": 3}  # the same questions are not added twice
    bad = admin.post("/api/admin/eval/cases/import", files={"file": ("set.csv", b"question,category\nWhat is X?,Podcast\n", "text/csv")})
    assert bad.status_code == 400 and "category" in bad.json()["detail"]
    cases = ok(admin.get("/api/admin/eval/cases"))
    row = next(c for c in cases if c["question"].startswith("How do you find the direction"))
    assert row["expected_locations"] == ["p. 2", "p. 1"] and row["subject_id"] == F["subject"]

    drafted = ok(admin.post("/api/admin/eval/cases/draft", json={"subject_id": F["subject"], "n": 2, "off_material": 2}))
    inm = [c for c in drafted if not c["off_material"]]
    assert inm and all(c["expected_locations"] and c["expected_answer"] and c["subject_id"] == F["subject"] for c in inm)
    assert len([c for c in drafted if c["off_material"]]) == 2
    assert admin.post("/api/admin/eval/cases/draft", json={"n": 99}).status_code == 422


def test_question_bank_preparation_for_the_simulation(admin):
    bank = ok(admin.get("/api/admin/eval/bank"))
    row = next(t for t in bank["topics"] if t["topic_id"] == F["topic"])
    assert bank["min"] == 8 and row["ready"] and row["verified"] >= 8  # topped up by the simulation test
    assert {"verified", "disagreement", "rejected"} <= set(bank["stats"])

    ok(admin.post("/api/admin/eval/bank/prepare", json={"per_topic": 20}))  # runs inline in tests
    bank = ok(admin.get("/api/admin/eval/bank"))
    assert bank["prep"]["status"] == "done" and bank["prep"]["done"] == bank["prep"]["total"] >= 1
    assert bank["prep"]["generated"] > 0 or bank["prep"]["skipped"]
    assert admin.post("/api/admin/eval/bank/prepare", json={"per_topic": 3}).status_code == 422


def test_a_crashed_bank_job_can_be_restarted(admin):
    from datetime import datetime, timedelta, timezone

    from app.db import SessionLocal
    from app.models import AppSetting
    from app.services import evaluation

    old = (datetime.now(timezone.utc) - timedelta(hours=2)).isoformat()
    with SessionLocal() as db:
        db.merge(AppSetting(key="bank_prep", value={"status": "running", "done": 1, "total": 5, "generated": 3,
                                                    "skipped": [], "updated_at": old}))
        db.commit()
        st = evaluation.bank_prep_state(db)
        assert st["status"] == "done" and st["stopped"]
    assert ok(admin.get("/api/admin/eval/bank"))["prep"]["stopped"]

    with SessionLocal() as db:
        db.merge(AppSetting(key="bank_prep", value={"status": "running", "done": 1, "total": 5, "generated": 3, "skipped": [],
                                                    "updated_at": datetime.now(timezone.utc).isoformat()}))
        db.commit()
    assert admin.post("/api/admin/eval/bank/prepare", json={"per_topic": 10}).status_code == 400
    with SessionLocal() as db:
        db.merge(AppSetting(key="bank_prep", value={"status": "done", "done": 0, "total": 0, "generated": 0, "skipped": []}))
        db.commit()


def test_evaluation_results_per_question_and_reports(admin, monkeypatch):
    from app.services import evaluation

    def partial_ragas(rows):  # a framework that couldn't score context recall
        return [{"faith": 0.9, "relev": 0.8, "cprec": 0.7, "crec": None} for _ in rows]

    monkeypatch.setitem(evaluation.FRAMEWORKS, "RAGAS", ("ragas", partial_ragas))
    run = ok(admin.post("/api/admin/eval/runs", json={"framework": "RAGAS"}))
    run = ok(admin.get(f"/api/admin/eval/runs/{run['id']}"))
    assert run["status"] == "completed", run["error"]
    m = run["metrics"]
    n_in = m["n_cases"] - m["n_off"]
    assert m["engine"] == "ragas" and m["values"]["faith"] == 0.9
    assert m["engine_note"].startswith(f"{n_in} of {4 * n_in} scores RAGAS couldn't produce")
    assert len(m["cases"]) == m["n_cases"]
    c = next(x for x in m["cases"] if x["q"].startswith("How do you find the direction"))
    assert c["crec"] is not None and c["faith"] == 0.9 and c["expected"] == ["p. 2", "p. 1"] and c["latency_s"] >= 0
    off = next(x for x in m["cases"] if x["q"] == "What is the capital of Australia?")
    assert off["off"] and off["faith"] is None
    ex = m["extra"]
    assert 0 <= ex["answered"] <= 1 and ex["latency_p90_s"] >= 0 and ex["models"]["num_ctx"] >= 8192

    md = admin.get(f"/api/admin/eval/runs/{run['id']}/report.md")
    assert md.status_code == 200 and md.headers["content-type"].startswith("text/markdown")
    assert f"run {run['number']}" in md.text and "## Per question" in md.text and "capital of Australia" in md.text
    assert f'intellinova-eval-run-{run["number"]}.md' in md.headers["content-disposition"]
    cs = admin.get(f"/api/admin/eval/runs/{run['id']}/report.csv")
    assert cs.status_code == 200 and cs.text.count("\n") == m["n_cases"] + 1

    stud = F["student"]
    assert stud.get(f"/api/admin/eval/runs/{run['id']}/report.md").status_code in (401, 403)
    assert stud.get("/api/admin/eval/bank").status_code in (401, 403)


def test_student_app_and_admin_console_are_separate(admin, stamp):
    from app.security import PORTAL_HEADER

    me = ok(admin.get("/api/auth/me"))
    assert me["is_staff"]
    # Admin accounts can't use the student app's API, and the web app is told to send them to the console.
    for path in ("/api/dashboard", "/api/me/header", "/api/learn/subjects", "/api/tutor/conversations", "/api/schedule"):
        r = admin.get(path)
        assert r.status_code == 403 and r.headers.get(PORTAL_HEADER) == "admin", path
    # ...nor log in on the student login page.
    c = new_client()
    r = c.post("/api/auth/login", json={"email": me["email"], "password": "adminpass1"})
    assert r.status_code == 403 and "/admin/login" in r.json()["detail"] and "inn_at" not in c.cookies
    assert ok(c.post("/api/auth/login", json={"email": me["email"], "password": "adminpass1", "portal": "admin"}))["user"]["is_staff"]

    # Students can't use the admin console, nor log in on the admin login page.
    st = F["student"]
    assert st.get("/api/admin/overview").status_code in (401, 403)
    assert st.get("/api/admin/eval/runs").status_code in (401, 403)
    tok = st.cookies.get("inn_at")
    assert st.get("/api/admin/overview", headers={"Authorization": f"Bearer {tok}"}).status_code == 403
    email = ok(st.get("/api/auth/me"))["email"]
    s2 = new_client()
    assert s2.post("/api/auth/login", json={"email": email, "password": "studentpw1", "portal": "admin"}).status_code == 403
    assert ok(s2.post("/api/auth/login", json={"email": email, "password": "studentpw1"}))["status"] == "signed_in"
    assert ok(s2.get("/api/dashboard"))


def test_a_queued_run_has_no_results_yet(admin):
    from app.db import SessionLocal
    from app.models import EvalRun
    from app.services import evaluation

    with SessionLocal() as db:
        r = EvalRun(number=evaluation.next_number(db), framework="TruLens", status="queued")
        db.add(r)
        db.commit()
        rid = str(r.id)
    row = ok(admin.get(f"/api/admin/eval/runs/{rid}"))
    assert row["metrics"] is None and row["simulation"] is None and row["question_bank"] is None and row["categories"] == []
    assert admin.get(f"/api/admin/eval/runs/{rid}/report.md").status_code == 200
    with SessionLocal() as db:
        db.delete(db.get(EvalRun, rid))
        db.commit()


def _mail_count(email: str) -> int:
    import httpx

    from tests.conftest import MAILPIT

    msgs = httpx.get(f"{MAILPIT}/api/v1/messages").json()["messages"]
    return sum(1 for m in msgs if any(t["Address"] == email for t in m["To"]))


def test_no_silent_missing_codes(admin, stamp):
    """Every path that promises an email either sends one or says why it can't."""
    import time

    # A fresh sign-up gets its code by email.
    c = new_client()
    email = f"neha{stamp}@example.com"
    assert ok(c.post("/api/auth/signup", json={"name": "Neha Gupta", "email": email, "password": "studentpw1",
                                               "accepted_terms": True}))["status"] == "verify"
    code = mail_code(email)
    assert len(code) == 6

    # Logging in before entering the code sends a new code and asks for it (instead of a dead end).
    time.sleep(61)  # GoTrue allows one email per address per minute
    n = _mail_count(email)
    assert ok(c.post("/api/auth/login", json={"email": email, "password": "studentpw1"}))["status"] == "verify"
    for _ in range(40):
        if _mail_count(email) > n:
            break
        time.sleep(0.25)
    assert _mail_count(email) > n
    ok(c.post("/api/auth/verify", json={"type": "signup", "email": email, "token": mail_code(email)}))

    # Signing up again with a confirmed address: GoTrue sends nothing, so the API says so.
    r = new_client().post("/api/auth/signup", json={"name": "Neha", "email": email, "password": "otherpass9", "accepted_terms": True})
    assert r.status_code == 409 and "already exists" in r.json()["detail"]
    admin_email = ok(admin.get("/api/auth/me"))["email"]
    r = new_client().post("/api/auth/signup", json={"name": "Asha", "email": admin_email, "password": "otherpass9", "accepted_terms": True})
    assert r.status_code == 409 and "/admin/login" in r.json()["detail"]

    # Password reset for an unknown address, or from the wrong portal, is refused instead of silently ignored.
    assert new_client().post("/api/auth/recover", json={"email": f"nobody{stamp}@example.com"}).status_code == 404
    assert new_client().post("/api/auth/recover", json={"email": email, "portal": "admin"}).status_code == 404
    assert new_client().post("/api/auth/recover", json={"email": admin_email}).status_code == 404
    assert ok(new_client().post("/api/auth/recover", json={"email": email}))["ok"]


def test_admin_and_student_signed_in_side_by_side(admin, stamp):
    """One browser (one cookie jar): an admin in the console and a student in the student app. Each app keeps
    its own session, so logging in or out in one never changes the other."""
    admin_email = ok(admin.get("/api/auth/me"))["email"]
    student_email = ok(F["student"].get("/api/auth/me"))["email"]
    A, S = {"X-IntelliNova-Portal": "admin"}, {"X-IntelliNova-Portal": "student"}

    b = new_client()
    ok(b.post("/api/auth/login", json={"email": admin_email, "password": "adminpass1", "portal": "admin"}, headers=A))
    assert {"inn_aat", "inn_art", "inn_as"} <= set(b.cookies.keys()) and "inn_at" not in b.cookies
    ok(b.post("/api/auth/login", json={"email": student_email, "password": "studentpw1"}, headers=S))
    assert {"inn_at", "inn_rt", "inn_s"} <= set(b.cookies.keys())

    # Both apps work at the same time, each as its own account.
    assert ok(b.get("/api/auth/me", headers=A))["email"] == admin_email
    assert ok(b.get("/api/auth/me", headers=S))["email"] == student_email
    assert ok(b.get("/api/admin/overview", headers=A))
    assert ok(b.get("/api/dashboard", headers=S))
    assert b.get("/api/admin/overview").status_code == 200  # /api/admin uses the console session by itself

    # Refreshing one session leaves the other alone.
    b.cookies.delete("inn_aat")
    assert ok(b.post("/api/auth/refresh", headers=A))["user"]["email"] == admin_email
    assert ok(b.get("/api/auth/me", headers=S))["email"] == student_email

    # Logging out of the student app keeps the admin console signed in.
    ok(b.post("/api/auth/logout", headers=S))
    assert "inn_at" not in b.cookies and "inn_s" not in b.cookies
    assert ok(b.get("/api/admin/overview", headers=A))
    assert b.get("/api/dashboard", headers=S).status_code == 401

    # A staff session left in the student app's cookies (from before the split) is refused and cleared on refresh.
    old = new_client()
    r = old.post("/api/auth/login", json={"email": admin_email, "password": "adminpass1", "portal": "admin"})
    old.cookies.set("inn_rt", old.cookies.get("inn_art"), path="/api/auth")
    assert old.post("/api/auth/refresh", headers=S).status_code == 401 and "inn_s" not in old.cookies
    assert r.status_code == 200


def test_citations_open_at_the_exact_place(admin, stamp):
    """Opening a citation's original: PDF pages and uploaded-video moments, from both apps."""
    import uuid

    from app.db import SessionLocal
    from app.models import ContentUnit, KbSource
    from app.services import storage

    path = storage.save_bytes(b"\x00\x00\x00\x18ftypmp42" + b"\x00" * 64, ".mp4", "sources")
    with SessionLocal() as db:
        src = KbSource(kind="video", title="Lecture 3", origin="admin", status="Ready", file_path=path, mime="video/mp4",
                       subject_id=uuid.UUID(F["subject"]))
        db.add(src)
        db.flush()
        u = ContentUnit(source_id=src.id, kind="Slide", text="Bulbs are filled with nitrogen and argon.", location="00:32",
                        t_start=32.0, position=0)
        db.add(u)
        db.commit()
        uid, sid = str(u.id), str(src.id)
    for client in (F["student"], admin):
        v = ok(client.get(f"/api/units/{uid}"))
        assert v["open_url"] == f"/api/sources/{sid}/file#t=32" and v["time"] == "00:32"
        r = client.get(f"/api/sources/{sid}/file")
        assert r.status_code == 200 and r.headers["content-type"].startswith("video/")


def test_tutor_answers_from_the_material_the_student_picks(admin, stamp):
    """Ask Tutor scope: all material, one subject, one chapter, my uploads, or particular files."""
    from sqlalchemy import select

    from app.db import SessionLocal
    from app.models import KbSource

    c = F["student"]
    opts = ok(c.get("/api/tutor/scopes"))
    sub = next(s for s in opts["subjects"] if s["id"] == F["subject"])
    assert any(ch["id"] == F["chapter"] for ch in sub["chapters"])
    magnet = next(s for s in opts["sources"] if s["title"] == "Magnetism notes")
    with SessionLocal() as db:
        other = db.scalar(select(KbSource).where(KbSource.origin == "admin", KbSource.id != magnet["id"],
                                                 KbSource.status.in_(("Ready", "Needs Review"))))
    assert other is not None
    q = {"content": "What do magnetic field lines never do?"}

    # A chapter: answered from that chapter's material; the choice is saved on the conversation.
    conv = ok(c.post("/api/tutor/conversations", json={"source_only": True, "scope": {"kind": "chapter", "id": F["chapter"]}}))
    assert conv["scope"]["kind"] == "chapter" and "Magnetic Effects of Current" in conv["scope"]["label"]
    r = ok(c.post(f"/api/tutor/conversations/{conv['id']}/messages", json=q))["assistant"]
    assert r["status"] == "answered" and r["searched"] == conv["scope"]["label"]
    assert all(ci["source_title"] == "Magnetism notes" for ci in r["citations"])

    # Switch mid-conversation to a file that doesn't cover it: declined, and the reply says what was searched.
    p = ok(c.patch(f"/api/tutor/conversations/{conv['id']}", json={"scope": {"kind": "sources", "ids": [str(other.id)]}}))
    assert p["scope"] == {"kind": "sources", "ids": [str(other.id)], "label": other.title}
    r = ok(c.post(f"/api/tutor/conversations/{conv['id']}/messages", json=q))["assistant"]
    assert r["status"] == "declined" and r["searched"] == other.title
    assert ok(c.get(f"/api/tutor/conversations/{conv['id']}"))["scope"]["kind"] == "sources"

    # Back to everything.
    p = ok(c.patch(f"/api/tutor/conversations/{conv['id']}", json={"scope": {"kind": "all"}}))
    assert p["scope"] == {"kind": "all", "label": "All my material"}
    r = ok(c.post(f"/api/tutor/conversations/{conv['id']}/messages", json=q))["assistant"]
    assert r["status"] == "answered" and r["searched"] is None

    # A subject works too.
    p = ok(c.patch(f"/api/tutor/conversations/{conv['id']}", json={"scope": {"kind": "subject", "id": F["subject"]}}))
    assert p["scope"]["label"] == sub["name"]

    # Invalid choices are refused: no uploads of their own yet, unknown ids, another student's private file.
    assert c.patch(f"/api/tutor/conversations/{conv['id']}", json={"scope": {"kind": "mine"}}).status_code in (200, 400)
    assert c.patch(f"/api/tutor/conversations/{conv['id']}", json={"scope": {"kind": "chapter", "id": "nope"}}).status_code == 400
    assert c.patch(f"/api/tutor/conversations/{conv['id']}", json={"scope": {"kind": "bogus"}}).status_code == 400
    with SessionLocal() as db:
        private = KbSource(kind="text", title=f"Someone else's notes {stamp}", origin="student", owner_id=None, status="Ready")
        db.add(private)
        db.commit()
        pid = str(private.id)
    r = c.patch(f"/api/tutor/conversations/{conv['id']}", json={"scope": {"kind": "sources", "ids": [pid]}})
    assert r.status_code == 400
    assert pid not in {s["id"] for s in ok(c.get("/api/tutor/scopes"))["sources"]}
